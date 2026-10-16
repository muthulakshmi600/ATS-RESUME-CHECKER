"""
Resume Parser Module for ResumeAI ATS.
Extracts clean plain text and structural cues from PDF and DOCX files.
Handles multi-page PDFs, scanned/image-only PDFs, and empty or corrupted documents.
"""

import os
import io
import re

ALLOWED_EXTENSIONS = {'pdf', 'docx'}
MAX_FILE_SIZE_BYTES = 5 * 1024 * 1024  # 5 MB

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def extract_text_from_pdf(file_stream_or_path):
    """
    Extracts text from PDF using PyMuPDF (fitz) or PyPDF2.
    Detects scanned image-only PDFs and empty files.
    """
    raw_bytes = None
    if isinstance(file_stream_or_path, (str, os.PathLike)):
        with open(file_stream_or_path, 'rb') as f:
            raw_bytes = f.read()
    elif isinstance(file_stream_or_path, bytes):
        raw_bytes = file_stream_or_path
    else:
        file_stream_or_path.seek(0)
        raw_bytes = file_stream_or_path.read()
        file_stream_or_path.seek(0)

    if not raw_bytes or len(raw_bytes) == 0:
        raise ValueError("The uploaded PDF file is empty (0 bytes). Please upload a valid resume.")

    text = ""
    page_count = 0
    has_images = False

    # Primary engine: PyMuPDF for highest fidelity
    try:
        import pymupdf
        doc = pymupdf.open(stream=raw_bytes, filetype="pdf")
        page_count = len(doc)
        for page in doc:
            page_text = page.get_text("text")
            if page_text:
                text += page_text + "\n"
            if len(page.get_images()) > 0:
                has_images = True
        doc.close()

        if text.strip():
            return clean_extracted_text(text)
        elif page_count > 0:
            if has_images:
                raise ValueError("This PDF appears to be a scanned image without selectable text. Please upload a text-based PDF or apply OCR first.")
            else:
                raise ValueError("The uploaded PDF is empty and contains no readable text. Please upload a valid resume.")
    except ValueError:
        raise
    except Exception as mupdf_err:
        pass

    # Fallback engine: PyPDF2
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(io.BytesIO(raw_bytes))
        page_count = len(reader.pages)
        for page in reader.pages:
            t = page.extract_text()
            if t:
                text += t + "\n"

        if text.strip():
            return clean_extracted_text(text)
        elif page_count > 0:
            raise ValueError("This PDF appears to be a scanned image or contains no selectable text. Please upload a text-based PDF or apply OCR first.")
        else:
            raise ValueError("The uploaded PDF has 0 pages and contains no text.")
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"Failed to parse PDF document (it may be corrupted or encrypted): {str(e)}")

def extract_text_from_docx(file_stream_or_path):
    """
    Extracts text from Microsoft Word (.docx) file using python-docx.
    Handles paragraphs and tables.
    """
    raw_bytes = None
    if isinstance(file_stream_or_path, (str, os.PathLike)):
        with open(file_stream_or_path, 'rb') as f:
            raw_bytes = f.read()
    elif isinstance(file_stream_or_path, bytes):
        raw_bytes = file_stream_or_path
    else:
        file_stream_or_path.seek(0)
        raw_bytes = file_stream_or_path.read()
        file_stream_or_path.seek(0)

    if not raw_bytes or len(raw_bytes) == 0:
        raise ValueError("The uploaded DOCX file is empty (0 bytes). Please upload a valid resume.")

    try:
        import docx
        doc = docx.Document(io.BytesIO(raw_bytes))

        paragraphs = []
        for p in doc.paragraphs:
            if p.text.strip():
                paragraphs.append(p.text.strip())

        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if row_cells:
                    paragraphs.append(" | ".join(row_cells))

        text = "\n".join(paragraphs)
        if not text.strip():
            raise ValueError("The uploaded DOCX document is empty and contains no readable text. Please upload a valid resume.")

        return clean_extracted_text(text)
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"Failed to parse DOCX document (it may be corrupted or not a valid Word document): {str(e)}")

def extract_resume_text(file_storage, filename):
    """
    Dispatcher to extract text based on extension.
    """
    if not filename or '.' not in filename:
        raise ValueError("Uploaded file has no extension. Please upload a .pdf or .docx file.")
    ext = filename.rsplit('.', 1)[1].lower()
    if ext == 'pdf':
        return extract_text_from_pdf(file_storage)
    elif ext == 'docx':
        return extract_text_from_docx(file_storage)
    else:
        raise ValueError(f"Unsupported file format '.{ext}'. Only PDF and DOCX files are supported.")

def clean_extracted_text(raw_text):
    """
    Cleans extracted text: fixes hyphenation, normalizes whitespace, retains paragraphs.
    """
    if not raw_text:
        return ""
    # Normalize unicode spaces
    text = raw_text.replace('\u200b', '').replace('\ufeff', '')
    # Replace carriage returns
    text = text.replace('\r\n', '\n').replace('\r', '\n')
    # Collapse 3+ consecutive newlines into 2
    text = re.sub(r'\n{3,}', '\n\n', text)
    # Fix broken hyphenated words at line ends (e.g. "experi-\nence" -> "experience")
    text = re.sub(r'(\w+)-\n(\w+)', r'\1\2', text)
    return text.strip()
